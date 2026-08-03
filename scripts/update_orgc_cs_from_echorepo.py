#!/usr/bin/env python3

from __future__ import annotations

import argparse
import math
import re
import shutil
from collections import defaultdict
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

IMAGE_EXTENSIONS = {
    ".jpg",
    ".jpeg",
    ".jfif",
    ".png",
    ".webp",
    ".tif",
    ".tiff",
}

def parse_sample_code(value: Any) -> tuple[str | None, str | None]:
    """
    Return:
        base_code: ABCD
        full_code: ABCD-1234, or None when only ABCD is available

    Examples:
        ABCD            -> ("ABCD", None)
        abcd            -> ("ABCD", None)
        ABCD-1234       -> ("ABCD", "ABCD-1234")
        ABCD1234        -> ("ABCD", "ABCD-1234")
        ECHO-ABCD-1234  -> ("ABCD", "ABCD-1234")
        sample ABCD     -> ("ABCD", None)
    """
    if value is None or pd.isna(value):
        return None, None

    text = str(value).strip().upper()

    if not text:
        return None, None

    text = re.sub(r"^ECHO[\s_-]*", "", text)

    full_match = re.search(
        r"(?<![A-Z0-9])([A-Z]{4})[\s_-]?(\d{4})(?!\d)",
        text,
    )

    if full_match:
        base = full_match.group(1)
        full = f"{base}-{full_match.group(2)}"
        return base, full

    short_match = re.search(
        r"(?<![A-Z0-9])([A-Z]{4})(?![A-Z0-9])",
        text,
    )

    if short_match:
        return short_match.group(1), None

    return None, None


def build_image_full_code_lookup(
    images_dir: Path,
) -> tuple[
    dict[str, str],
    dict[str, list[str]],
    dict[str, list[str]],
]:
    """
    Inspect image filenames and map four-letter base codes to full codes.

    Examples:
        WVSU-1234.jpg  -> WVSU -> WVSU-1234
        AABB-5678.jfif -> AABB -> AABB-5678

    Returns:
        unique_lookup:
            base code -> one unambiguous full code

        ambiguous_lookup:
            base code -> multiple different full codes

        image_paths:
            full code -> corresponding image paths
    """
    unique_lookup: dict[str, str] = {}
    ambiguous_lookup: dict[str, list[str]] = {}
    image_paths: dict[str, list[str]] = defaultdict(list)

    if not images_dir.exists():
        print(f"WARNING: image directory not found: {images_dir}")
        return unique_lookup, ambiguous_lookup, image_paths

    full_codes_by_base: dict[str, set[str]] = defaultdict(set)

    for path in sorted(images_dir.rglob("*")):
        if not path.is_file():
            continue

        if path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue

        base, full = parse_sample_code(path.stem)

        # Filenames such as ABCD.jpg cannot resolve a conflict.
        if base is None or full is None:
            continue

        full_codes_by_base[base].add(full)
        image_paths[full].append(path.as_posix())

    for base, full_codes in full_codes_by_base.items():
        sorted_codes = sorted(full_codes)

        if len(sorted_codes) == 1:
            unique_lookup[base] = sorted_codes[0]
        else:
            ambiguous_lookup[base] = sorted_codes

    return unique_lookup, ambiguous_lookup, image_paths


def parse_percent_number(value: Any) -> float | None:
    """
    Parse citizen values such as:
        2.5%  -> 2.5
        2,5%  -> 2.5
        2.5   -> 2.5
    """
    if value is None or pd.isna(value):
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else None

    text = str(value).strip()

    if not text:
        return None

    if text.lower() in {"nan", "none", "null", "na", "n/a"}:
        return None

    text = text.replace("\u00a0", "")
    text = text.replace(" ", "")
    text = text.replace("%", "")
    text = text.replace(",", ".")

    try:
        number = float(text)
    except ValueError:
        return None

    return number if math.isfinite(number) else None


def find_column(
    columns: list[str],
    requested: str,
) -> str:
    """Find a column case-insensitively."""
    lower_lookup = {
        str(column).strip().lower(): column
        for column in columns
    }

    result = lower_lookup.get(requested.strip().lower())

    if result is None:
        raise ValueError(
            f"Column {requested!r} was not found.\n"
            f"Available columns:\n  "
            + "\n  ".join(str(column) for column in columns)
        )

    return result


def find_excel_column(
    worksheet,
    header_row: int,
    requested: str,
) -> int | None:
    """Find an Excel column number case-insensitively."""
    requested = requested.strip().lower()

    for cell in worksheet[header_row]:
        if cell.value is None:
            continue

        if str(cell.value).strip().lower() == requested:
            return cell.column

    return None


def unique_numbers(values: list[float], tolerance: float = 1e-9) -> list[float]:
    unique: list[float] = []

    for value in values:
        if not any(abs(value - existing) <= tolerance for existing in unique):
            unique.append(value)

    return unique


def resolve_exact_groups(
    groups: dict[str, list[float]],
    key_to_base: dict[str, str],
    policy: str,
) -> tuple[dict[str, float], list[dict[str, Any]]]:
    """
    Resolve repeated occurrences of the same exact source identifier.

    These are true duplicate-source conflicts, unlike different full QR codes
    that happen to share the same four-letter base.
    """
    lookup: dict[str, float] = {}
    conflicts: list[dict[str, Any]] = []

    for exact_key, values in groups.items():
        distinct = unique_numbers(values)

        if len(distinct) == 1:
            lookup[exact_key] = distinct[0]
            continue

        conflicts.append(
            {
                "base_code": key_to_base[exact_key],
                "exact_source_code": exact_key,
                "all_values": "|".join(str(v) for v in values),
                "distinct_values": "|".join(str(v) for v in distinct),
            }
        )

        if policy == "first":
            lookup[exact_key] = values[0]
        elif policy == "last":
            lookup[exact_key] = values[-1]
        elif policy == "mean":
            lookup[exact_key] = sum(values) / len(values)
        elif policy == "error":
            continue
        else:
            raise ValueError(f"Unsupported conflict policy: {policy}")

    return lookup, conflicts


def build_base_fallback_lookup(
    exact_lookup: dict[str, float],
    key_to_base: dict[str, str],
) -> tuple[dict[str, float], dict[str, list[tuple[str, float]]]]:
    """
    Build a fallback lookup for four-letter IDs.

    A base code is safe only when all matching source records have the same
    SOIL_COLOR_color value.

    Example safe fallback:
        AABB-1234 -> 2.5
        AABB-5678 -> 2.5

    Example ambiguous fallback:
        WVSU-1234 -> 3.5
        WVSU-5678 -> 1.5
    """
    entries_by_base: dict[str, list[tuple[str, float]]] = defaultdict(list)

    for exact_key, value in exact_lookup.items():
        base = key_to_base[exact_key]
        entries_by_base[base].append((exact_key, value))

    base_lookup: dict[str, float] = {}
    ambiguous: dict[str, list[tuple[str, float]]] = {}

    for base, entries in entries_by_base.items():
        values = [value for _, value in entries]
        distinct = unique_numbers(values)

        if len(distinct) == 1:
            base_lookup[base] = distinct[0]
        else:
            ambiguous[base] = entries

    return base_lookup, ambiguous


def report_path(output: Path, suffix: str) -> Path:
    return output.with_name(f"{output.stem}_{suffix}.csv")


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Update orgC_CS in a laboratory workbook using "
            "ECHOREPO SOIL_COLOR_color values."
        )
    )
    parser.add_argument(
        "--images-dir",
        default="data/samples/with_gray",
        help=(
            "Directory containing images whose filenames may provide complete "
            "sample codes such as ABCD-1234.jpg."
        ),
    )
    parser.add_argument(
        "--lab",
        default="data/lab/test_stat_orgC.xlsx",
    )
    parser.add_argument(
        "--source",
        default="data/context/echorepo_samples_with_email.csv",
    )
    parser.add_argument(
        "--output",
        default="data/lab/test_stat_orgC_with_orgC_CS.xlsx",
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Worksheet name. The active sheet is used by default.",
    )
    parser.add_argument(
        "--header-row",
        type=int,
        default=1,
    )
    parser.add_argument(
        "--lab-id-col",
        default="ID",
    )
    parser.add_argument(
        "--source-id-col",
        default="QR_qrCode",
    )
    parser.add_argument(
        "--source-value-col",
        default="SOIL_COLOR_color",
    )
    parser.add_argument(
        "--target-col",
        default="orgC_CS",
    )
    parser.add_argument(
        "--exact-conflict-policy",
        choices=["error", "first", "last", "mean"],
        default="error",
        help=(
            "How to handle different values repeated for the same exact "
            "source QR code. This does not affect ambiguous four-letter codes."
        ),
    )
    parser.add_argument(
        "--clear-unmatched",
        action="store_true",
    )

    args = parser.parse_args()

    lab_path = Path(args.lab)
    source_path = Path(args.source)
    output_path = Path(args.output)

    if not lab_path.exists():
        raise FileNotFoundError(f"Lab workbook not found: {lab_path}")

    if not source_path.exists():
        raise FileNotFoundError(f"Source CSV not found: {source_path}")

    source = pd.read_csv(
        source_path,
        sep=None,
        engine="python",
        dtype=str,
        keep_default_na=False,
        encoding="utf-8-sig",
    )

    source.columns = [str(column).strip() for column in source.columns]

    source_id_col = find_column(
        list(source.columns),
        args.source_id_col,
    )
    source_value_col = find_column(
        list(source.columns),
        args.source_value_col,
    )

    print(f"Source ID column: {source_id_col}")
    print(f"Source value column: {source_value_col}")

    exact_groups: dict[str, list[float]] = defaultdict(list)
    key_to_base: dict[str, str] = {}
    invalid_source_rows: list[dict[str, Any]] = []

    for source_index, row in source.iterrows():
        raw_code = row[source_id_col]
        raw_value = row[source_value_col]

        base, full = parse_sample_code(raw_code)
        value = parse_percent_number(raw_value)

        if base is None:
            invalid_source_rows.append(
                {
                    "source_row": source_index + 2,
                    "raw_code": raw_code,
                    "raw_value": raw_value,
                    "reason": "invalid_sample_code",
                }
            )
            continue

        if value is None:
            invalid_source_rows.append(
                {
                    "source_row": source_index + 2,
                    "raw_code": raw_code,
                    "raw_value": raw_value,
                    "reason": "missing_or_invalid_SOIL_COLOR_color",
                }
            )
            continue

        # Preserve the complete QR code whenever available.
        exact_key = full if full is not None else base

        exact_groups[exact_key].append(value)
        key_to_base[exact_key] = base

    exact_lookup, exact_conflicts = resolve_exact_groups(
        exact_groups,
        key_to_base,
        policy=args.exact_conflict_policy,
    )

    exact_conflict_file = report_path(output_path, "exact_source_conflicts")
    exact_conflict_file.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(exact_conflicts).to_csv(
        exact_conflict_file,
        index=False,
        encoding="utf-8-sig",
    )

    if exact_conflicts and args.exact_conflict_policy == "error":
        print("")
        print(
            "Different values remain for the same exact QR code. "
            "See:"
        )
        print(f"  {exact_conflict_file}")
        print("")
        print("First conflicts:")

        for conflict in exact_conflicts[:20]:
            print(
                f"  {conflict['exact_source_code']}: "
                f"{conflict['distinct_values']}"
            )

        raise SystemExit(
            "Resolve these exact-code conflicts, or explicitly choose "
            "--exact-conflict-policy first, last, or mean."
        )

    base_lookup, ambiguous_bases = build_base_fallback_lookup(
        exact_lookup,
        key_to_base,
    )

    images_dir = Path(args.images_dir)

    (
        image_full_code_lookup,
        ambiguous_image_codes,
        image_paths,
    ) = build_image_full_code_lookup(images_dir)

    print("")
    print(f"Images directory: {images_dir}")
    print(
        "Base codes resolved from full image filenames: "
        f"{len(image_full_code_lookup)}"
    )
    print(
        "Base codes with multiple full image filenames: "
        f"{len(ambiguous_image_codes)}"
    )
    ambiguous_rows = []

    for base, entries in sorted(ambiguous_bases.items()):
        ambiguous_rows.append(
            {
                "base_code": base,
                "source_codes": "|".join(code for code, _ in entries),
                "values": "|".join(str(value) for _, value in entries),
            }
        )

    ambiguous_file = report_path(output_path, "ambiguous_base_codes")
    pd.DataFrame(ambiguous_rows).to_csv(
        ambiguous_file,
        index=False,
        encoding="utf-8-sig",
    )

    workbook = load_workbook(lab_path)

    if args.sheet:
        if args.sheet not in workbook.sheetnames:
            raise ValueError(
                f"Sheet {args.sheet!r} not found. "
                f"Available sheets: {workbook.sheetnames}"
            )

        worksheet = workbook[args.sheet]
    else:
        worksheet = workbook.active

    lab_id_col = find_excel_column(
        worksheet,
        args.header_row,
        args.lab_id_col,
    )

    if lab_id_col is None:
        raise ValueError(
            f"Excel column {args.lab_id_col!r} was not found "
            f"in sheet {worksheet.title!r}."
        )

    target_col = find_excel_column(
        worksheet,
        args.header_row,
        args.target_col,
    )

    if target_col is None:
        target_col = worksheet.max_column + 1
        worksheet.cell(
            row=args.header_row,
            column=target_col,
            value=args.target_col,
        )
        print(f"Created target column: {args.target_col}")

    match_report: list[dict[str, Any]] = []

    exact_matches = 0
    image_matches = 0
    base_matches = 0
    ambiguous_matches = 0
    ambiguous_image_matches = 0
    unmatched = 0
    invalid_lab_ids = 0

    for row_number in range(args.header_row + 1, worksheet.max_row + 1):
        id_cell = worksheet.cell(row=row_number, column=lab_id_col)
        target_cell = worksheet.cell(row=row_number, column=target_col)

        raw_id = id_cell.value

        if raw_id is None or str(raw_id).strip() == "":
            continue

        base, full = parse_sample_code(raw_id)
        old_value = target_cell.value

        status: str
        new_value: float | None = None
        matched_source_key: str | None = None

        image_full_code: str | None = None

        if base is None:
            status = "invalid_lab_id"
            invalid_lab_ids += 1

        # Best case: Excel already contains the complete code.
        elif full is not None and full in exact_lookup:
            new_value = exact_lookup[full]
            matched_source_key = full
            status = "matched_exact_full_code"
            exact_matches += 1

        # Excel has only ABCD, but the renamed image provides ABCD-1234.
        elif (
            full is None
            and base in image_full_code_lookup
        ):
            image_full_code = image_full_code_lookup[base]

            if image_full_code in exact_lookup:
                new_value = exact_lookup[image_full_code]
                matched_source_key = image_full_code
                status = "matched_via_image_filename"
                image_matches += 1
            else:
                status = "image_full_code_not_found_in_source"
                unmatched += 1

        # No useful full image filename, but every CSV record under this
        # base code has the same value.
        elif base in base_lookup:
            new_value = base_lookup[base]
            matched_source_key = base
            status = "matched_unambiguous_base_fallback"
            base_matches += 1

        # More than one different full code exists among image filenames.
        elif base in ambiguous_image_codes:
            status = "multiple_full_image_codes_not_updated"
            ambiguous_image_matches += 1

        # CSV contains different values under this four-letter prefix,
        # and no renamed image resolved which full code is correct.
        elif base in ambiguous_bases:
            status = "ambiguous_base_not_updated"
            ambiguous_matches += 1

        else:
            status = "unmatched"
            unmatched += 1

        if new_value is not None:
            target_cell.value = new_value
        elif args.clear_unmatched:
            target_cell.value = None

        match_report.append(
            {
                "excel_row": row_number,
                "ID": raw_id,
                "base_code": base or "",
                "full_code": full or "",
                "image_full_code": image_full_code or "",
                "image_paths": (
                    "|".join(image_paths.get(image_full_code, []))
                    if image_full_code
                    else ""
                ),
                "matched_source_key": matched_source_key or "",
                "status": status,
                "old_orgC_CS": old_value,
                "new_orgC_CS": target_cell.value,
            }
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)

    if output_path.resolve() == lab_path.resolve():
        backup_path = lab_path.with_suffix(".before_orgC_CS_update.xlsx")
        shutil.copy2(lab_path, backup_path)
        print(f"Backup created: {backup_path}")

    workbook.save(output_path)

    match_file = report_path(output_path, "match_report")
    invalid_file = report_path(output_path, "invalid_source_rows")

    pd.DataFrame(match_report).to_csv(
        match_file,
        index=False,
        encoding="utf-8-sig",
    )

    pd.DataFrame(invalid_source_rows).to_csv(
        invalid_file,
        index=False,
        encoding="utf-8-sig",
    )

    print("")
    print("Update completed")
    print(f"Output workbook: {output_path}")
    print(f"Exact full-code matches: {exact_matches}")
    print(f"Matches resolved through image filename: {image_matches}")
    print(
        "Rows with multiple full image codes: "
        f"{ambiguous_image_matches}"
    )
    print(f"Unambiguous base-code matches: {base_matches}")
    print(f"Ambiguous base codes not updated: {ambiguous_matches}")
    print(f"Unmatched Excel IDs: {unmatched}")
    print(f"Invalid Excel IDs: {invalid_lab_ids}")
    print(f"Exact source conflicts: {len(exact_conflicts)}")
    print("")
    print(f"Match report: {match_file}")
    print(f"Ambiguous base-code report: {ambiguous_file}")
    print(f"Exact source-conflict report: {exact_conflict_file}")
    print(f"Invalid source rows: {invalid_file}")


if __name__ == "__main__":
    main()